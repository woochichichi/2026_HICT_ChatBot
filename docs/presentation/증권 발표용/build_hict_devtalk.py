# -*- coding: utf-8 -*-
"""
증권/ITO AX Day 발표덱 — v3 다크 dev-talk 스타일 (전면 리디자인)
피드백: 원본 30분 덱 양식과 완전히 다르게, 일반적인 IT 기술 공유회 느낌, 필수 내용만.
 - 다크 배경 + 단일 민트 액센트 + 모노 태그 (개발자 기술 공유회 톤)
 - 카드/다크밴드 모티프 폐기 → 타이포 중심 미니멀
 - 13장으로 축약 (필수만)
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- 다크 테마 토큰 ----
BG    = RGBColor(0x15, 0x17, 0x1C)   # near-black
PANEL = RGBColor(0x1E, 0x21, 0x2A)   # subtle raised surface
PANEL2= RGBColor(0x27, 0x2B, 0x36)
TX    = RGBColor(0xEC, 0xEC, 0xEE)   # main text
MUT   = RGBColor(0x9A, 0x9E, 0xAB)   # muted
DIM   = RGBColor(0x54, 0x58, 0x66)   # very dim (inactive nav)
MINT  = RGBColor(0x5E, 0xEA, 0xD4)   # accent
MINTD = RGBColor(0x2D, 0xD4, 0xBF)
LINE  = RGBColor(0x30, 0x34, 0x40)   # hairline
KR    = "맑은 고딕"
MONO  = "Consolas"

SECTIONS = ["배경", "만든 것", "시연", "기술", "도입·비전"]

prs = Presentation()
prs.slide_width  = Emu(12192000)
prs.slide_height = Emu(6858000)
BLANK = prs.slide_layouts[6]

def _font(run, size, color, bold, font=KR):
    run.font.name = font; run.font.size = Pt(size); run.font.color.rgb = color; run.font.bold = bold
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {}); rPr.append(ea)
    ea.set('typeface', font)

def txt(s, l, t, w, h, text, size=12, color=TX, bold=False,
        align=PP_ALIGN.LEFT, font=KR, anchor=MSO_ANCHOR.TOP, wrap=True, line=None, sp=None):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = wrap; tf.vertical_anchor = anchor
    for m in ('margin_left','margin_right','margin_top','margin_bottom'): setattr(tf, m, 0)
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line is not None: p.line_spacing = line
        if sp is not None: p.space_after = Pt(sp)
        runs = ln if isinstance(ln, list) else [(ln, {})]
        for seg, ov in runs:
            r = p.add_run(); r.text = seg
            _font(r, ov.get('size', size), ov.get('color', color), ov.get('bold', bold), ov.get('font', KR))
    return tb

def rect(s, l, t, w, h, fill, line_c=None, radius=None):
    shp_t = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = s.shapes.add_shape(shp_t, Inches(l), Inches(t), Inches(w), Inches(h))
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line_c is None: shp.line.fill.background()
    else: shp.line.color.rgb = line_c; shp.line.width = Pt(1)
    shp.shadow.inherit = False
    if radius:
        try: shp.adjustments[0] = radius
        except Exception: pass
    return shp

def newslide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = BG; bg.line.fill.background(); bg.shadow.inherit=False
    return s

def _label_w(name):
    kw = sum(1.02 if ord(c) > 0x1100 else 0.6 for c in name)  # 한글 vs 라틴/공백/·
    return (3*0.6 + kw) * 10.5/72 + 0.12   # "0N " 접두 + 라벨, 여유 0.12in

def navbar(s, active):
    """미니멀 상단 진행 표시 — 활성=민트, 지난=MUT, 남은=DIM. pill/막대 없음."""
    x = 0.9; y = 0.5
    for i, name in enumerate(SECTIONS):
        if i == active: col, bold = MINT, True
        elif active is not None and i < active: col, bold = MUT, False
        else: col, bold = DIM, False
        w = _label_w(name)
        txt(s, x, y, w, 0.28, f"0{i+1} {name}", 10.5, col, bold, font=MONO, wrap=False)
        x += w + 0.4
    txt(s, 9.9, 0.5, 2.53, 0.28, "ITO · AX DAY", 10.5, DIM, False, font=MONO, align=PP_ALIGN.RIGHT, wrap=False)

def head(s, active, kicker, title, subtitle=None):
    navbar(s, active)
    txt(s, 0.9, 1.28, 11.5, 0.3, kicker, 12, MINT, True, font=MONO)
    txt(s, 0.9, 1.66, 11.6, 0.7, title, 27, TX, True)
    if subtitle:
        txt(s, 0.9, 2.42, 11.6, 0.4, subtitle, 13, MUT)

def take(s, text):
    """하단 한 줄 테이크어웨이 (다크밴드 대신)."""
    txt(s, 0.9, 6.72, 11.6, 0.4, [[("→  ", {'color':MINT,'bold':True,'size':13}), (text, {'color':TX,'size':13,'bold':True})]], 13, TX)

def pageno(s, n):
    txt(s, 11.5, 6.95, 0.9, 0.2, f"{n:02d}", 9, DIM, font=MONO, align=PP_ALIGN.RIGHT)

# =====================================================================
# 1. TITLE
# =====================================================================
s = newslide()
txt(s, 0.9, 1.7, 11.0, 0.3, "2026 ITO AX DAY", 13, MINT, True, font=MONO)
txt(s, 0.9, 2.25, 11.4, 1.5,
    [[("편람을 답하는 AI를", {'color':TX})], [("직접 ", {'color':TX}), ("만들었습니다", {'color':MINT})]],
    44, TX, True, line=1.05)
txt(s, 0.9, 4.25, 11.2, 0.5, "RAG 상담 어시스턴트 — 만든 것 · 넘은 벽 · 앞으로의 계획", 15, MUT)
# 모노 태그
txt(s, 0.9, 5.15, 11.0, 0.3, "#RAG   #폐쇄망_온프레미스   #하이브리드검색   #사내전환", 12, MINTD, False, font=MONO)
# 하단 발표자/날짜
rect(s, 0.9, 6.35, 3.0, 0.02, LINE)
txt(s, 0.9, 6.5, 8.0, 0.3, "ITO팀  ·  2026.07", 12, MUT, font=MONO)
txt(s, 0.9, 6.9, 11.4, 0.28, "RAG: AI가 답을 지어내지 않고 사내 문서에서 근거를 먼저 찾아 답하는 방식", 9, DIM)

# =====================================================================
# 2. TL;DR
# =====================================================================
s = newslide()
navbar(s, None)
txt(s, 0.9, 1.28, 11.5, 0.3, "TL;DR", 12, MINT, True, font=MONO)
txt(s, 0.9, 1.66, 11.6, 0.7, "세 줄 요약", 27, TX, True)
tl = [
    ("만들었다", "편람을 출처와 함께 답하는 RAG를 팀이 직접 구현하고 수치로 검증"),
    ("넘었다", "폐쇄망·인증·최신성·외부 AI — 현장의 벽 네 가지를 직접 해결"),
    ("계획한다", "사내 AI 전환과 확장까지 설계 — 끌고 갈 역량을 보여준다"),
]
y = 2.9
for i, (h, d) in enumerate(tl):
    txt(s, 0.9, y, 1.2, 0.9, f"0{i+1}", 34, MINT, True, font=MONO)
    txt(s, 2.25, y+0.02, 3.0, 0.5, h, 21, TX, True)
    txt(s, 5.4, y+0.1, 7.0, 0.6, d, 14, MUT, line=1.25)
    if i < 2: rect(s, 0.9, y+1.05, 11.5, 0.012, LINE)
    y += 1.28
take(s, "제품을 파는 자리가 아니라, 이만큼 만드는 팀이라는 걸 보여주는 자리")
pageno(s, 2)

# =====================================================================
# 3. [배경]
# =====================================================================
s = newslide()
head(s, 0, "01 · 왜 만들었나", "상담의 병목은 사람이 아니라 '흩어진 지식'",
     "제도·상품이 자주 바뀌는 증권 상담 — 시간을 잡아먹는 건 '지식을 찾는 방식'")
probs = [
    ("분산된 지식", "편람·위키·공지·메신저가 흩어져 검색 경로가 길다"),
    ("변경 반영 시차", "제도·상품 변경 후 현장 숙지까지 시차 → 오안내 위험"),
    ("응대 편차", "숙련도에 따라 검색어·표현·누락이 달라 품질이 들쭉날쭉"),
]
y = 3.15
for i, (h, d) in enumerate(probs):
    txt(s, 0.9, y, 0.7, 0.4, f"0{i+1}", 15, MINT, True, font=MONO)
    txt(s, 1.7, y-0.02, 3.4, 0.4, h, 16, TX, True)
    txt(s, 5.2, y+0.02, 7.2, 0.4, d, 13, MUT)
    if i < 2: rect(s, 0.9, y+0.5, 11.5, 0.012, LINE)
    y += 0.75
# 영상 자리 (borderless, mint 좌측 마커)
rect(s, 0.9, 5.55, 0.06, 0.55, MINT)
txt(s, 1.15, 5.55, 11.2, 0.55, [[("▶ 상담 현장 영상  ", {'color':MINT,'bold':True,'size':13,'font':MONO}),
    ("지금 상담원이 겪는 상황을 먼저 보여준 뒤, 우리가 만든 화면으로", {'color':MUT,'size':12})]],
    12, MUT, anchor=MSO_ANCHOR.MIDDLE)
take(s, "문제는 상담원이 아니라, 지식을 찾는 '방식'이다")
pageno(s, 3)

# =====================================================================
# 4. [만든 것]
# =====================================================================
s = newslide()
head(s, 1, "02 · 무엇을 만들었나", "세 가지 일을 하나의 도구로",
     "단일 데모가 아니라 상담 현장에 필요한 세 업무를 통합")
sol = [
    ("답변 검색", "질문에 필요한 답·근거를 즉시", "+ 추천 응대 멘트 · 출처 링크"),
    ("업무편람 RAG", "절차·예외·처리 방법을 출처와 함께", "+ 문서·페이지 단위 근거"),
    ("신입 교육 (AI 코치)", "AI 고객과 대화 → 즉시 채점", "+ 모범답안 · 페르소나 8종"),
]
cw, gap, x0, y = 3.73, 0.3, 0.9, 3.15
for i, (h, d, tag) in enumerate(sol):
    x = x0 + i*(cw+gap)
    rect(s, x, y, cw, 2.5, PANEL, radius=0.05)
    txt(s, x+0.32, y+0.32, 0.9, 0.4, f"0{i+1}", 15, MINT, True, font=MONO)
    txt(s, x+0.32, y+0.85, cw-0.64, 0.5, h, 16.5, TX, True)
    txt(s, x+0.32, y+1.4, cw-0.64, 0.55, d, 12.5, MUT, line=1.3)
    txt(s, x+0.32, y+2.02, cw-0.64, 0.35, tag, 11, MINTD, False, font=MONO)
take(s, "처음엔 상담원 보조 도구로. 검증되면 접점을 넓힌다")
pageno(s, 4)

# =====================================================================
# 5. [시연]
# =====================================================================
s = newslide()
head(s, 2, "03 · 시연", "답을 찾는 AI, 사람을 키우는 AI",
     "개념이 아니라 실제로 동작하는 화면 (기존 5분 발표 영상 활용)")
demo = [
    ("응대 모드", "상담원 Copilot", ["두서없는 질문도 근거를 먼저 검색", "정리된 답변 + 추천 응대 멘트"], "출처 100% · 신뢰도 상/중/하"),
    ("AI 코치 모드", "즉시 채점", ["AI가 8종 고객으로 질문", "응대 → 점수·누락·모범답안 즉시"], "페르소나 8 · 난이도 3 · 약점 복습"),
]
cw, gap, x0, y = 5.75, 0.3, 0.9, 3.05
for i, (tag, h, ds, foot) in enumerate(demo):
    x = x0 + i*(cw+gap)
    rect(s, x, y, cw, 2.6, PANEL, radius=0.05)
    txt(s, x+0.35, y+0.32, cw-0.7, 0.3, tag, 11.5, MINT, True, font=MONO)
    txt(s, x+0.35, y+0.72, cw-0.7, 0.4, h, 18, TX, True)
    txt(s, x+0.35, y+1.3, cw-0.7, 0.7,
        [[("· ", {'color':MINT,'bold':True}), (d, {'color':MUT})] for d in ds], 12.5, MUT, line=1.4)
    txt(s, x+0.35, y+2.15, cw-0.7, 0.35, foot, 10.5, MINTD, font=MONO)
take(s, "편람을 찾아줄 뿐 아니라 응대 멘트까지 — 신입은 선배 없이 숙련된다")
pageno(s, 5)

# =====================================================================
# 6. [기술] 동작 방식 — 파이프라인
# =====================================================================
s = newslide()
head(s, 3, "04 · 어떻게 동작하나", "편람을 'AI가 찾을 수 있는 형태'로",
     "통째로 넣지 않고 '찾아지는 카드 묶음'으로 바꿔 둔다")
steps = [("수집", "위키·문서 원문"), ("분할", "의미 단위 조각"), ("색인", "뜻으로 검색"), ("저장", "출처와 함께")]
cw, gap, x0, y = 2.55, 0.45, 0.9, 3.3
for i, (h, d) in enumerate(steps):
    x = x0 + i*(cw+gap)
    rect(s, x, y, cw, 1.5, PANEL, radius=0.06)
    txt(s, x+0.3, y+0.28, cw-0.6, 0.3, f"0{i+1}", 12, MINT, True, font=MONO)
    txt(s, x+0.3, y+0.62, cw-0.6, 0.4, h, 17, TX, True)
    txt(s, x+0.3, y+1.08, cw-0.6, 0.35, d, 11.5, MUT)
    if i < 3:
        txt(s, x+cw+0.04, y+0.45, 0.4, 0.6, "→", 18, MINT, True, align=PP_ALIGN.CENTER)
txt(s, 0.9, 5.25, 11.5, 0.5, [[("여기에 ", {'color':MUT,'size':13}),
    ("답변 정확도를 높이는 장치", {'color':MINT,'bold':True,'size':13}),
    ("를 얹었다 — 다음 장에서.", {'color':MUT,'size':13})]], 13, MUT)
take(s, "편람 = '검색되는 카드 묶음' — 조각 + 의미 + 출처")
pageno(s, 6)

# =====================================================================
# 7. [기술] 어려웠던 점 4가지 개요
# =====================================================================
s = newslide()
head(s, 3, "05 · 어려웠던 점", "부딪힌 벽 4가지, 그리고 해결",
     "폐쇄망 금융사라면 똑같이 만나는 난점 — 하나씩 넘었다")
walls = [
    ("닫힌 회사 환경", "개발도구·자동 로그인 없음", "수집·가공 머신 분리 + Windows 무인 인증"),
    ("매일 바뀌는 편람", "매번 전체 재처리는 감당 불가", "'지문'으로 바뀐 문서만 골라 자동 갱신"),
    ("정확한 답과 근거", "엉뚱한 답·종목코드 누락", "뜻+단어 하이브리드 검색 · 출처 부착"),
    ("외부 AI 의존", "사내 AI 구동 환경 부재", "모델 교체 계층 — 사내 도입 시 바로 전환"),
]
y = 2.95
for i, (h, p, sol) in enumerate(walls):
    txt(s, 0.9, y, 0.7, 0.4, f"0{i+1}", 15, MINT, True, font=MONO)
    txt(s, 1.7, y-0.02, 3.5, 0.4, h, 15.5, TX, True)
    txt(s, 5.3, y+0.03, 3.3, 0.4, p, 12, MUT)
    txt(s, 8.7, y+0.03, 3.7, 0.4, [[("→ ", {'color':MINT,'bold':True}), (sol, {'color':MINT})]], 12, MINT)
    if i < 3: rect(s, 0.9, y+0.52, 11.5, 0.012, LINE)
    y += 0.78
take(s, "다음 두 장에서 가장 흥미로운 해결 과정을 자세히")
pageno(s, 7)

# =====================================================================
# 8. [기술] 딥다이브 ① 데이터 확보
# =====================================================================
def deepdive(page, kicker, title, subtitle, blocks, note, take_txt):
    s = newslide()
    head(s, 3, kicker, title, subtitle)
    cw, gap, x0, y, ch = 5.75, 0.3, 0.9, 3.0, 2.35
    for i, (p, sol) in enumerate(blocks):
        x = x0 + i*(cw+gap)
        rect(s, x, y, cw, ch, PANEL, radius=0.05)
        txt(s, x+0.35, y+0.3, cw-0.7, 0.65,
            [[("문제  ", {'color':MUT,'bold':True,'size':10.5,'font':MONO}), (p, {'color':TX,'bold':True,'size':13})]], 13, TX, line=1.25)
        txt(s, x+0.35, y+1.15, cw-0.7, ch-1.3,
            [[("해결  ", {'color':MINT,'bold':True,'size':10.5,'font':MONO}), (sol, {'color':MUT,'size':11.5})]], 11.5, MUT, line=1.4)
    if note:
        txt(s, 0.9, 5.6, 11.5, 0.4, [[("↳ ", {'color':MINT,'bold':True}), (note, {'color':MUT})]], 11.5, MUT)
    take(s, take_txt)
    pageno(s, page)

deepdive(8, "05 · 딥다이브 ①", "닫힌 폐쇄망에서 데이터 가져오기",
    "핵심 질문: 사내에만 있는 편람을 어떻게 가져오나",
    [
     ("사내망 PC엔 개발 도구를 깔 수 없다",
      "'꺼내는 일'과 '가공하는 일'을 다른 컴퓨터로 분리 — 금고에선 서류만 꺼내고 정리는 밖에서"),
     ("AI가 위키에 매번 사람이 로그인해줘야 한다",
      "로그인된 Windows 계정을 신분증처럼 활용 → 무인 접속. 매일 바뀌는 인증번호도 무관"),
    ],
    "화면에 접혀 있던 페이지까지 빠짐없이 수집 — 편람 누락 0건",
    "고객사도 똑같은 폐쇄망 — 이미 넘어온 길")

# =====================================================================
# 9. [기술] 딥다이브 ② 정확하고 믿을 수 있는 답
# =====================================================================
deepdive(9, "05 · 딥다이브 ②", "정확한 답, 그리고 '왜 그 답인지'",
    "발표에서 가장 궁금한 지점 — 정확한가, 근거를 댈 수 있는가",
    [
     ("'K-OTC' 같은 코드를 뜻으로만 찾으면 놓친다",
      "'뜻으로 찾기' + '단어로 찾기'를 겹친 하이브리드 검색 → 검색 정확도가 눈에 띄게 상승"),
     ("AI가 가끔 그럴듯한 엉뚱한 답을 낸다",
      "모든 답에 근거 문서·페이지 부착 → 답보다 '근거'를 먼저 확인. 틀려도 즉시 추적"),
    ],
    "평가에서 정답이 항상 상위 3건 안에 포함 확인 (문항 확대 검증 중)",
    "근거 없는 답은 내지 않는다 — 최종 판단은 사람")

# =====================================================================
# 10. [도입·비전] 실도입 시 고려
# =====================================================================
s = newslide()
head(s, 4, "06 · 실제 도입한다면", "이렇게 올린다 — 고려사항",
     "지금 당장은 아니어도, 결정되면 바로 올릴 수 있는 그림")
layers = [("화면","React UI"),("서비스","API·권한"),("검색·답변","RAG"),("지식 저장소","기업용 벡터DB"),("AI 모델","사내 전환")]
cw, gap, x0, y = 2.15, 0.24, 0.9, 3.05
for i, (h, d) in enumerate(layers):
    x = x0 + i*(cw+gap)
    hi = i >= 3
    rect(s, x, y, cw, 1.6, PANEL2 if hi else PANEL, radius=0.06)
    txt(s, x+0.2, y+0.32, cw-0.4, 0.4, h, 13.5, (MINT if hi else TX), True, align=PP_ALIGN.CENTER)
    txt(s, x+0.2, y+0.9, cw-0.4, 0.35, d, 10.5, MUT, align=PP_ALIGN.CENTER, font=MONO)
    if i < 4:
        txt(s, x+cw-0.02, y+0.5, 0.28, 0.5, "→", 15, MINT, True, align=PP_ALIGN.CENTER)
txt(s, 0.9, 5.15, 11.5, 0.5, [[("PoC는 가벼운 도구로 검증. ", {'color':MUT,'size':12.5}),
    ("실서비스는 별도 서버 + 기업용 벡터DB + 사내 GPU", {'color':MINT,'bold':True,'size':12.5}),
    (".", {'color':MUT})]], 12.5, MUT)
take(s, "계층을 분리해 — 나중에 '답변 모델'만 사내용으로 갈아끼운다")
pageno(s, 10)

# =====================================================================
# 11. [도입·비전] 앞으로
# =====================================================================
s = newslide()
head(s, 4, "07 · 앞으로", "확장의 청사진",
     "더 개발하겠다가 아니라, 도입 시 무엇이 필요한지를 제시")
# 좌: 확장 3단계
txt(s, 0.9, 2.95, 5.4, 0.3, "적용 범위", 12, MINT, True, font=MONO)
stages = [("지금", "업무편람 — 구현·검증 완료 ✅"), ("다음", "사내 문서 통합 (게시판·파일서버)"), ("비전", "개인별 지식 허브 (메일·회의록)")]
y = 3.4
for i, (h, d) in enumerate(stages):
    txt(s, 0.9, y, 1.4, 0.35, h, 14, TX, True)
    txt(s, 2.5, y+0.03, 3.9, 0.4, d, 12, MUT)
    if i < 2: rect(s, 0.9, y+0.48, 5.4, 0.012, LINE)
    y += 0.72
# 우: 필요한 것
txt(s, 6.9, 2.95, 5.5, 0.3, "필요한 것", 12, MINT, True, font=MONO)
needs = [("사내 AI 환경", "AI 구동 GPU 서버"), ("기업용 벡터DB", "운영용 저장소"), ("국산 금융 LLM", "데이터 주권·규제 대응")]
y = 3.4
for i, (h, d) in enumerate(needs):
    txt(s, 6.9, y, 2.4, 0.35, h, 14, TX, True)
    txt(s, 9.4, y+0.03, 3.0, 0.4, d, 12, MUT)
    if i < 2: rect(s, 6.9, y+0.48, 5.5, 0.012, LINE)
    y += 0.72
txt(s, 0.9, 5.85, 11.5, 0.3, "타 증권사도 이미 같은 방향으로 가고 있다.", 12, MINTD, font=MONO)
take(s, "이 프로젝트를 끌고 갈 역량이 있다 — 그것을 오늘 보여줬다")
pageno(s, 11)

# =====================================================================
# 12. Q&A
# =====================================================================
s = newslide()
txt(s, 0.9, 2.7, 11.0, 1.3, "Q&A", 60, MINT, True, font=MONO)
txt(s, 0.95, 4.1, 11.0, 0.5, "Thank you.", 22, TX, False)
txt(s, 0.95, 4.75, 11.0, 0.3, "ITO팀 · 2026 AX Day", 12, MUT, font=MONO)

# =====================================================================
# 13. 백업 (질문 대비 수치)
# =====================================================================
s = newslide()
txt(s, 0.9, 1.3, 11.5, 0.3, "APPENDIX · 질문 대비", 12, MINT, True, font=MONO)
txt(s, 0.9, 1.68, 11.6, 0.6, "정량 근거 수치", 25, TX, True)
txt(s, 0.9, 2.4, 11.6, 0.4, "본편에선 내세우지 않되, 물어보면 바로 답할 수 있는 근거", 13, MUT)
stats = [("편람 수집","520p","누락 0건"),("적재 조각","1,281","의미 단위"),("검색 정확도","top-3","정답 항상 포함"),("출처 부착","100%","모든 답변"),("재처리","변경분만","자동 갱신")]
cw, gap, x0, y = 2.15, 0.24, 0.9, 3.25
for i, (h, v, d) in enumerate(stats):
    x = x0 + i*(cw+gap)
    rect(s, x, y, cw, 1.75, PANEL, radius=0.06)
    txt(s, x+0.2, y+0.28, cw-0.4, 0.3, h, 10.5, MUT, True, align=PP_ALIGN.CENTER, font=MONO)
    txt(s, x+0.2, y+0.68, cw-0.4, 0.5, v, 21, MINT, True, align=PP_ALIGN.CENTER, font=MONO)
    txt(s, x+0.2, y+1.28, cw-0.4, 0.35, d, 10.5, MUT, align=PP_ALIGN.CENTER)
txt(s, 0.9, 5.35, 11.6, 0.6,
    "※ 답변 생성 정확도는 30문항 기준 — 문항 확대 검증 중. 억원·ROI 등 비용 수치는 산정 근거와 함께 별도 보관.",
    10.5, DIM, line=1.3)
take(s, "수치는 '증명'의 근거일 뿐 — 메시지는 '이만큼 만드는 팀'")
pageno(s, 13)

out = "/home/user/2026_HICT_ChatBot/docs/presentation/증권 발표용/HICT_PPT_15min_devtalk.pptx"
prs.save(out)
print("saved:", out)
print("slides:", len(prs.slides._sldIdLst))
