# -*- coding: utf-8 -*-
"""
증권 발표용 신규 PPT 빌드 (녹취록 260723 피드백 반영, 15분 분량)
- 녹취록 분석 결과를 반영해 기존 HICT_PPT.pptx(27장, 30분)를 재구성
- 디자인 토큰은 원본 덱과 동일 (한화 오렌지 FF6600 / 다크 2A2A2A / 맑은 고딕)
- 핵심 방향: '영업'이 아닌 '역량 증명(테크데이)', 3축=사례·역량·계획, 영업성 수치 제거
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- 디자인 토큰 (원본 덱에서 추출) ----
ORANGE = RGBColor(0xFF, 0x66, 0x00)
DARK   = RGBColor(0x2A, 0x2A, 0x2A)
GRAY   = RGBColor(0x56, 0x56, 0x56)
LGRAY  = RGBColor(0x8C, 0x8C, 0x89)
BORDER = RGBColor(0xD6, 0xD6, 0xD2)
CREAM  = RGBColor(0xFF, 0xF6, 0xEF)
BROWN  = RGBColor(0x7A, 0x5B, 0x42)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
BAND   = RGBColor(0x2A, 0x2A, 0x2A)
FAINT  = RGBColor(0xD9, 0xD9, 0xD6)
FONT   = "맑은 고딕"

prs = Presentation()
prs.slide_width  = Emu(12192000)   # 13.33in
prs.slide_height = Emu(6858000)    # 7.5in
BLANK = prs.slide_layouts[6]

def slide():
    return prs.slides.add_slide(BLANK)

def _set_font(run, size, color, bold, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    # 한글 폰트 지정 (동아시아)
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', font)

def txt(s, l, t, w, h, text, size=12, color=DARK, bold=False,
        align=PP_ALIGN.LEFT, font=FONT, anchor=MSO_ANCHOR.TOP, wrap=True, line=None):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    for m in ('margin_left','margin_right','margin_top','margin_bottom'):
        setattr(tf, m, 0)
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line is not None:
            p.line_spacing = line
        runs = ln if isinstance(ln, list) else [(ln, {})]
        for seg, ov in runs:
            r = p.add_run(); r.text = seg
            _set_font(r, ov.get('size', size), ov.get('color', color),
                      ov.get('bold', bold), ov.get('font', font))
    return tb

def card(s, l, t, w, h, fill=WHITE, linecol=BORDER, radius=0.06):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = linecol; shp.line.width = Pt(1)
    shp.shadow.inherit = False
    try:
        shp.adjustments[0] = radius
    except Exception:
        pass
    return shp

def band(s, text, l=0.84, t=6.05, w=11.65, h=0.62, size=11.5):
    b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    b.fill.solid(); b.fill.fore_color.rgb = BAND; b.line.fill.background()
    b.shadow.inherit = False
    tf = b.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.25); tf.margin_right = Inches(0.25)
    tf.margin_top = 0; tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    r = p.add_run(); r.text = text; _set_font(r, size, WHITE, True)
    return b

def header(s, kicker, num, title, subtitle, title_size=24):
    txt(s, 1.54, 0.46, 8.0, 0.2, kicker, 8.5, GRAY)
    if num:
        txt(s, 1.02, 0.66, 0.62, 0.55, num, 30, ORANGE, True, wrap=False)
    txt(s, 1.54, 0.66, 10.4, 0.6, title, title_size, DARK, True)
    if subtitle:
        txt(s, 1.54, 1.34, 10.6, 0.4, subtitle, 12.5, GRAY)

def footer(s, page):
    txt(s, 0.75, 7.12, 3.2, 0.18, "ⓒ Hanwha Systems. All rights reserved", 7.5, LGRAY)
    txt(s, 6.24, 7.12, 0.85, 0.18, f"- {page} -", 7.5, LGRAY, align=PP_ALIGN.CENTER)

KICK = "한화시스템 ITO · 2026 AX Day"

# =====================================================================
# SLIDE 1 — 표지
# =====================================================================
s = slide()
# 상단 얇은 라벨
txt(s, 1.1, 1.5, 11.0, 0.3, KICK, 13, ORANGE, True)
txt(s, 1.1, 2.05, 11.3, 1.4,
    [[("ITO ", {'color':DARK}), ("AX Day", {'color':ORANGE})]],
    38, DARK, True)
txt(s, 1.1, 3.15, 11.2, 0.7,
    "직접 만든 RAG 상담 어시스턴트 — 우리가 만든 것, 넘은 벽, 그리고 앞으로의 계획",
    15, GRAY)
# 3 키워드 카드 (짧게)
kw = [
    ("직접 만들었다", "업무편람을 출처와 함께 답하는 AI 상담 어시스턴트를 팀이 직접 구현·검증했습니다.", WHITE),
    ("폐쇄망을 넘었다", "개발 도구도 자동 로그인도 없던 회사 환경의 벽을 하나씩 풀어냈습니다.", CREAM),
    ("다음을 계획한다", "지금은 안 되는 것도, 사내 환경이 갖춰지면 바로 적용되도록 설계해 두었습니다.", WHITE),
]
cw, gap, x0, y = 3.55, 0.4, 1.1, 4.35
for i, (h, d, f) in enumerate(kw):
    x = x0 + i*(cw+gap)
    card(s, x, y, cw, 1.5, fill=f)
    txt(s, x+0.25, y+0.22, cw-0.5, 0.3, h, 14, ORANGE, True)
    txt(s, x+0.25, y+0.62, cw-0.5, 0.75, d, 9.5, GRAY, line=1.15)
txt(s, 1.1, 6.2, 6.0, 0.25, "2026.07", 11, GRAY, True)
txt(s, 1.1, 6.6, 11.2, 0.3,
    "RAG: AI가 답을 지어내지 않고 사내 문서에서 근거를 먼저 찾아 답하는 방식 · AX: AI 기반 업무 전환",
    8.2, LGRAY)

# =====================================================================
# SLIDE 2 — 오늘 세 가지 (사례·역량·계획)
# =====================================================================
s = slide()
header(s, KICK, "", "오늘, 이 세 가지만 기억하시면 됩니다", "사례 → 역량 → 계획. 제품을 팔러 온 자리가 아니라, 이만큼 만드는 팀이라는 걸 보여드립니다.", 26)
cols = [
    ("①  사례", "직접 만든 RAG", "업무편람 520페이지를 수집해, AI가 출처와 함께 답하도록 팀이 직접 만들고 눈으로 확인했습니다.", WHITE),
    ("②  역량", "폐쇄망의 벽을 넘은 실행력", "개발 도구 없는 PC, 위키 인증, 매일 바뀌는 편람 — 현장의 벽을 피하지 않고 직접 풀었습니다.", CREAM),
    ("③  계획", "실도입·확장의 청사진", "지금 못 하는 부분도 사내 환경이 갖춰지면 바로 올라가도록 설계했고, 다음 확장까지 그렸습니다.", WHITE),
]
cw, gap, x0, y = 3.72, 0.35, 1.1, 2.1
for i, (tag, h, d, f) in enumerate(cols):
    x = x0 + i*(cw+gap)
    card(s, x, y, cw, 3.4, fill=f)
    txt(s, x+0.3, y+0.32, cw-0.6, 0.3, tag, 15, ORANGE, True)
    txt(s, x+0.3, y+0.95, cw-0.6, 0.9, h, 18, DARK, True, line=1.1)
    txt(s, x+0.3, y+2.0, cw-0.6, 1.2, d, 11, GRAY, line=1.3)
band(s, "제품을 파는 자리가 아닙니다. ITO팀이 이만큼 만들 수 있는 팀이라는 걸 보여드리는 자리입니다.")
footer(s, 2)

# =====================================================================
# SLIDE 3 — 배경
# =====================================================================
s = slide()
header(s, KICK, "01", "상담의 병목은 사람이 아니라 '흩어진 지식'입니다", "제도·상품이 자주 바뀌는 증권 상담에서, 정작 시간을 잡아먹는 건 '지식을 찾는 방식'입니다.")
probs = [
    ("분산된 지식", "업무편람·위키·공지·메신저가 여기저기 흩어져 있어, 필요한 내용을 찾는 경로가 깁니다."),
    ("변경 반영 시차", "제도·상품·수수료가 바뀌어도 현장이 숙지하기까지 시간이 걸려 오안내 위험이 생깁니다."),
    ("응대 편차", "상담원 숙련도에 따라 검색어·답변 표현·누락 항목이 달라져 품질이 들쭉날쭉합니다."),
]
cw, gap, x0, y = 3.72, 0.35, 1.1, 2.1
for i, (h, d) in enumerate(probs):
    x = x0 + i*(cw+gap)
    card(s, x, y, cw, 2.35)
    txt(s, x+0.3, y+0.3, 0.5, 0.4, "0"+str(i+1), 20, ORANGE, True)
    txt(s, x+0.3, y+0.85, cw-0.6, 0.35, h, 15, DARK, True)
    txt(s, x+0.3, y+1.3, cw-0.6, 0.9, d, 11, GRAY, line=1.3)
# 배경 영상 안내 (녹취록: 상담 현장 영상을 먼저 보여준다)
vb = card(s, 1.1, 4.62, 11.4, 0.62, fill=CREAM)
txt(s, 1.35, 4.72, 11.0, 0.42, [[("▶  상담 현장 영상  ", {'color':ORANGE,'bold':True,'size':12}),
    ("지금 상담원이 겪는 상황을 먼저 보여드린 뒤, 우리가 만든 화면으로 이어집니다.", {'color':GRAY,'size':11})]],
    11, GRAY, anchor=MSO_ANCHOR.MIDDLE)
band(s, "문제는 상담원이 아니라, 지식을 찾는 '방식'입니다.")
footer(s, 3)

# =====================================================================
# SLIDE 4 — 솔루션 개요
# =====================================================================
s = slide()
header(s, KICK, "02", "우리가 만든 것 — 세 가지 일을 하나로", "단일 데모가 아니라, 상담 현장에서 실제로 필요한 세 가지 업무를 하나의 도구로 묶었습니다.")
sol = [
    ("답변 검색", "고객 질문에 필요한 답과 근거를 즉시 찾아, 추천 응대 멘트와 출처까지 함께 제시합니다."),
    ("업무편람 RAG", "업무 절차·예외 기준·처리 방법을, 어느 문서 몇 페이지인지 출처와 함께 알려줍니다."),
    ("신입 교육 (AI 코치)", "AI가 고객이 되어 질문하면 신입이 응대하고, 즉시 채점·모범답안을 받습니다. (페르소나 8종)"),
]
cw, gap, x0, y = 3.72, 0.35, 1.1, 2.1
for i, (h, d) in enumerate(sol):
    x = x0 + i*(cw+gap)
    card(s, x, y, cw, 2.9)
    txt(s, x+0.3, y+0.3, 0.5, 0.4, "0"+str(i+1), 20, ORANGE, True)
    txt(s, x+0.3, y+0.88, cw-0.6, 0.35, h, 15.5, DARK, True)
    txt(s, x+0.3, y+1.35, cw-0.6, 1.3, d, 11, GRAY, line=1.35)
band(s, "처음엔 상담원의 내부 보조 도구로. 성능·보안·운영이 검증되면 접점을 넓혀갑니다.")
footer(s, 4)

# =====================================================================
# SLIDE 5 — 시연
# =====================================================================
s = slide()
header(s, KICK, "03", "시연 — 답을 찾는 AI, 사람을 키우는 AI", "개념이 아니라 실제로 동작하는 두 개의 화면을 보여드립니다.")
demo = [
    ("응대 모드 — 상담원 Copilot", "두서없는 고객 질문도 편람 근거를 먼저 찾아, 정리된 답변과 추천 응대 멘트까지 제시합니다.",
     "모든 답변에 출처 부착 · 신뢰도 상/중/하 표시"),
    ("AI 코치 모드 — 즉시 채점", "AI가 8종 고객 페르소나로 질문하면 신입이 응대를 작성하고, 점수·누락 항목·모범답안을 즉시 받습니다.",
     "페르소나 8종 · 난이도 3단계 · 약점 자동 복습"),
]
cw, gap, x0, y = 5.6, 0.45, 1.1, 2.1
for i, (h, d, tag) in enumerate(demo):
    x = x0 + i*(cw+gap)
    card(s, x, y, cw, 2.55, fill=(CREAM if i==0 else WHITE))
    txt(s, x+0.32, y+0.3, cw-0.64, 0.35, h, 15.5, DARK, True)
    txt(s, x+0.32, y+0.85, cw-0.64, 1.1, d, 11.5, GRAY, line=1.4)
    txt(s, x+0.32, y+1.95, cw-0.64, 0.4, tag, 10.5, ORANGE, True)
vb = card(s, 1.1, 4.85, 11.4, 0.55, fill=WHITE)
txt(s, 1.35, 4.9, 11.0, 0.45, [[("▶  우리가 만든 화면  ", {'color':ORANGE,'bold':True,'size':12}),
    ("실제 질문을 던지면 근거와 함께 답이 나오는 과정을 영상으로 보여드립니다.", {'color':GRAY,'size':11})]],
    11, GRAY, anchor=MSO_ANCHOR.MIDDLE)
band(s, "편람을 찾아줄 뿐 아니라 응대 멘트까지 — 신입은 선배 없이도 반복 훈련으로 숙련됩니다.")
footer(s, 5)

# =====================================================================
# SLIDE 6 — 데이터 전처리 (간단히)
# =====================================================================
s = slide()
header(s, KICK, "03", "편람을 'AI가 찾을 수 있는 형태'로 바꿉니다",
       "편람을 통째로 넣지 않습니다. '찾아지는 카드 묶음'으로 바꿔 두는 것이 핵심입니다.")
steps = [
    ("위키·문서 수집", "흩어진 원문을 모읍니다"),
    ("조각으로 나누기", "의미 단위로 잘라둡니다"),
    ("의미로 색인", "AI가 '뜻'으로 찾게 합니다"),
    ("출처와 함께 저장", "어디서 왔는지 붙여 둡니다"),
]
cw, gap, x0, y = 2.55, 0.45, 1.15, 2.35
for i, (h, d) in enumerate(steps):
    x = x0 + i*(cw+gap)
    card(s, x, y, cw, 1.9)
    txt(s, x+0.25, y+0.28, 0.6, 0.4, "0"+str(i+1), 18, ORANGE, True)
    txt(s, x+0.25, y+0.82, cw-0.5, 0.35, h, 13.5, DARK, True)
    txt(s, x+0.25, y+1.22, cw-0.5, 0.5, d, 10, GRAY, line=1.25)
    if i < 3:
        txt(s, x+cw+0.02, y+0.6, 0.42, 0.6, "→", 20, LGRAY, True, align=PP_ALIGN.CENTER)
txt(s, 1.15, 4.55, 11.3, 0.5,
    [[("여기에 더해, ", {'color':GRAY,'size':12}),
      ("답변 정확도를 높이는 우리만의 장치", {'color':ORANGE,'bold':True,'size':12}),
      ("를 얹었습니다. (뒤 '어려웠던 점'에서 설명)", {'color':GRAY,'size':12})]],
    12, GRAY)
band(s, "편람을 통째로 저장하지 않고, '검색되는 카드 묶음'으로 바꿔 둡니다.")
footer(s, 6)

# =====================================================================
# SLIDE 7 — 어려웠던 점 (개요)
# =====================================================================
s = slide()
header(s, KICK, "04", "만들면서 부딪힌 '벽', 그리고 넘은 방법",
       "편람을 AI가 읽게 만드는 길에는 벽이 있었습니다 — 고객사 환경에도 똑같이 서 있는 벽입니다.")
walls = [
    ("①", "닫힌 회사 환경", "개발 도구도, 자동 로그인도 없던 회사 PC에서 데이터를 어떻게 가져오나"),
    ("②", "매일 바뀌는 편람", "조금씩 바뀌는 내용을, 매번 전부 다시 하지 않고 어떻게 따라잡나"),
    ("③", "정확한 답과 근거", "엉뚱한 답을 막고, 종목코드까지 정확히 찾아, '왜 그 답인지'를 붙이기"),
    ("④", "외부 AI 의존", "지금은 외부 AI를 쓰지만, 사내 AI로 갈아끼울 수 있게 준비하기"),
]
cw, gap, x0, y = 2.78, 0.24, 0.94, 2.15
for i, (n, h, d) in enumerate(walls):
    x = x0 + i*(cw+gap)
    card(s, x, y, cw, 3.25, fill=(CREAM if i==3 else WHITE))
    txt(s, x+0.28, y+0.3, 0.6, 0.35, n, 16, ORANGE, True)
    txt(s, x+0.28, y+0.85, cw-0.56, 0.7, h, 14.5, DARK, True, line=1.1)
    txt(s, x+0.28, y+1.7, cw-0.56, 1.4, d, 10.5, GRAY, line=1.35)
band(s, "모두 폐쇄망 금융사라면 똑같이 만나는 벽입니다. 우리는 이미 하나씩 넘어봤습니다.")
footer(s, 7)

# =====================================================================
# 벽 상세 공통 렌더러 (문제 → 접근/해결)
# =====================================================================
def wall_detail(page, num, title, subtitle, blocks, note=None, band_txt=""):
    s = slide()
    header(s, KICK, "04", title, subtitle)
    n = len(blocks)
    cw = 5.6 if n == 2 else 11.4
    gap = 0.45
    x0, y = 1.1, 2.05
    ch = 2.7 if n == 2 else 1.4
    for i, blk in enumerate(blocks):
        if n == 2:
            x = x0 + i*(cw+gap); yy = y
        else:
            x = x0; yy = y + i*(ch+0.25)
        prob, sol = blk
        card(s, x, yy, cw, ch)
        txt(s, x+0.32, yy+0.26, cw-0.64, 0.55,
            [[("문제  ", {'color':ORANGE,'bold':True,'size':11}), (prob, {'color':DARK,'bold':True,'size':12.5})]],
            12.5, DARK, line=1.2)
        ty = yy + (1.15 if n==2 else 0.72)
        txt(s, x+0.32, ty, cw-0.64, ch-1.2,
            [[("해결  ", {'color':RGBColor(0x1E,0x7A,0x3C),'bold':True,'size':11}), (sol, {'color':GRAY,'size':11})]],
            11, GRAY, line=1.35)
    if note:
        txt(s, 1.1, 5.55, 11.4, 0.4, note, 10.5, LGRAY, line=1.25)
    band(s, band_txt)
    footer(s, page)
    return s

# SLIDE 8 — 벽 ① 닫힌 환경 (수집·인증 병합)
wall_detail(8, "04", "벽 ① — 닫힌 회사 환경에서 데이터를 가져오기",
    "핵심 질문은 하나였습니다. '사내에만 있는 업무편람 데이터를 어떻게 가져오나?'",
    [
     ("사내망에서만 위키에 닿는데, 그 PC엔 개발 도구를 깔 수 없었습니다.",
      "문서를 '꺼내는 일'과 '가공하는 일'을 다른 컴퓨터로 분리했습니다. 금고에서 서류만 꺼내고, 정리는 밖에서 하는 방식입니다."),
     ("AI가 위키에 접근하려면 매번 사람이 로그인을 해줘야 했습니다.",
      "이미 로그인된 Windows 계정을 신분증처럼 활용해, 사람 개입 없이 접속하게 했습니다. 인증번호가 매일 바뀌어도 문제없습니다."),
    ],
    note="화면에 접혀 있던 페이지까지 빠짐없이 모아, 편람을 누락 없이 확보했습니다.",
    band_txt="고객사도 똑같은 폐쇄망 환경입니다. 우리는 이미 넘어온 길입니다.")

# SLIDE 9 — 벽 ② 매일 바뀌는 편람
wall_detail(9, "04", "벽 ② — 매일 조금씩 바뀌는 편람 따라잡기",
    "편람은 매일 조금씩 바뀝니다. 바뀔 때마다 전부 다시 처리하면 시간도 비용도 감당이 안 됩니다.",
    [
     ("바뀐 부분이 어딘지 모르면, 매번 편람 전체를 처음부터 다시 처리해야 합니다.",
      "문서마다 '지문(내용 표시)'을 찍어두고, 바뀐 문서만 골라 다시 처리합니다. 오타 한 줄 고쳤다고 책 전체를 재인쇄하지 않고, 그 페이지만 교체하는 방식입니다."),
     ("데이터가 바뀌어도 상담원이 늘 최신 내용을 봐야 신뢰할 수 있습니다.",
      "매일 자동으로 변경분만 반영해, 불필요한 재처리 없이 항상 최신 상태를 유지합니다."),
    ],
    band_txt="규정이 수시로 바뀌는 증권 업무에 맞는 방식입니다. 바뀐 만큼만, 자동으로.")

# SLIDE 10 — 벽 ③ 정확한 답과 근거
wall_detail(10, "04", "벽 ③ — 정확한 답, 그리고 '왜 그 답인지'",
    "발표에서 가장 궁금해하실 지점입니다. 답이 정확한가, 그리고 그 근거를 댈 수 있는가.",
    [
     ("'K-OTC' 같은 종목코드·고유명사를, 뜻으로만 찾으면 놓칩니다.",
      "'뜻으로 찾기'와 '단어로 찾기'를 함께 씁니다. 두 방식을 겹치자 검색 정확도가 눈에 띄게 올랐습니다."),
     ("AI가 가끔 그럴듯하지만 엉뚱한 답을 내놓습니다.",
      "모든 답에 근거 문서·페이지를 붙였습니다. 상담원은 답보다 '근거'를 먼저 보고 판단합니다. 틀려도 어디서 왔는지 즉시 확인됩니다."),
    ],
    note="평가에서 정답 문서가 항상 상위 3건 안에 포함됨을 확인했고, 문항을 넓혀 검증을 이어가고 있습니다.",
    band_txt="근거 없는 답은 내지 않습니다. AI가 아니라 상담원이 최종 판단합니다.")

# SLIDE 11 — 벽 ④ 외부 AI → 사내 AI
wall_detail(11, "04", "벽 ④ — 지금은 외부 AI, 그러나 사내 AI로 갈 준비",
    "사내에서 AI 모델을 돌릴 환경이 아직 없어, 지금은 외부 AI를 씁니다. 이걸 어떻게 안전하게, 그리고 나중에 어떻게 바꾸나.",
    [
     ("사내에 AI를 돌릴 환경(GPU 등)이 아직 없어 외부 AI에 의존합니다.",
      "지금도 고객 데이터는 쓰지 않고, 밖으로 나가는 데이터는 기록으로 남깁니다. 문서를 '찾는' 부분은 이미 사내용 모델로 전환을 마쳤습니다."),
     ("나중에 사내 AI가 도입되면, 시스템을 다시 만들어야 하지 않나?",
      "'모델만 갈아끼우면' 되도록 교체 계층을 미리 설계했습니다. 사내망에 AI가 들어오는 순간 바로 그것으로 바꿀 수 있습니다."),
    ],
    band_txt="외부에 의존하는 지금도, 사내로 들어갈 다음도 — 둘 다 준비돼 있습니다.")

# =====================================================================
# SLIDE 12 — 실도입 청사진 (아키텍처 + DB)
# =====================================================================
s = slide()
header(s, KICK, "05", "실제 도입한다면 — 이렇게 올립니다",
       "지금 당장은 아니어도, 도입이 결정되면 바로 올릴 수 있는 그림을 준비했습니다.")
layers = [
    ("상담원 화면", "질문을 넣고\n답·근거를 봅니다"),
    ("서비스(API)", "권한·세션을\n관리합니다"),
    ("검색·답변(RAG)", "근거를 찾아\n답을 구성합니다"),
    ("지식 저장소", "기업용 벡터 DB에\n안전하게 보관"),
    ("AI 모델", "사내 모델로\n갈아끼웁니다"),
]
cw, gap, x0, y = 2.15, 0.28, 1.1, 2.3
for i, (h, d) in enumerate(layers):
    x = x0 + i*(cw+gap)
    fill = CREAM if i in (3,4) else WHITE
    card(s, x, y, cw, 1.95, fill=fill)
    txt(s, x+0.2, y+0.3, cw-0.4, 0.7, h, 13, DARK, True, align=PP_ALIGN.CENTER, line=1.05)
    txt(s, x+0.2, y+1.05, cw-0.4, 0.75, d, 9.5, GRAY, align=PP_ALIGN.CENTER, line=1.2)
    if i < 4:
        txt(s, x+cw-0.02, y+0.7, 0.32, 0.5, "→", 16, LGRAY, True, align=PP_ALIGN.CENTER)
txt(s, 1.1, 4.55, 11.4, 0.5,
    [[("PoC에선 가벼운 도구로 빠르게 검증했습니다. ", {'color':GRAY,'size':11.5}),
      ("실서비스에선 별도 서버와 기업용 벡터 DB", {'color':ORANGE,'bold':True,'size':11.5}),
      ("로 올려 안정성과 보안을 확보합니다.", {'color':GRAY,'size':11.5})]],
    11.5, GRAY, line=1.3)
band(s, "화면·서비스·검색·저장소·모델을 분리해, 나중에 '답변 모델'만 사내용으로 갈아끼울 수 있습니다.")
footer(s, 12)

# =====================================================================
# SLIDE 13 — 기대효과 (정성 중심, 영업성 수치 제거)
# =====================================================================
s = slide()
header(s, KICK, "05", "이걸로 무엇이 좋아지나 — 세 가지",
       "숫자로 겁주는 자리가 아닙니다. 무엇이 어떻게 나아지는지를 먼저 말씀드립니다.")
eff = [
    ("응대 시간이 줄어듭니다", "흩어진 지식을 찾아 헤매는 시간이 줄어, 고객 대기와 상담원 부담이 함께 줄어듭니다."),
    ("답변 품질이 고르게 표준화됩니다", "누가 응대해도 같은 근거로 답하게 되어, 숙련도에 따른 편차와 오안내 위험이 줄어듭니다."),
    ("신입 교육이 빨라집니다", "AI 코치로 반복 훈련해, 선배가 붙지 않아도 신입이 빠르게 숙련됩니다."),
]
cw, gap, x0, y = 3.72, 0.35, 1.1, 2.15
for i, (h, d) in enumerate(eff):
    x = x0 + i*(cw+gap)
    card(s, x, y, cw, 2.55, fill=(CREAM if i==1 else WHITE))
    txt(s, x+0.3, y+0.3, 0.5, 0.4, "0"+str(i+1), 20, ORANGE, True)
    txt(s, x+0.3, y+0.9, cw-0.6, 0.7, h, 14.5, DARK, True, line=1.1)
    txt(s, x+0.3, y+1.65, cw-0.6, 0.8, d, 11, GRAY, line=1.35)
txt(s, 1.1, 4.9, 11.4, 0.4,
    "정량 효과(절감 시간·비용 등)는 별도 근거 자료로 준비돼 있으며, 필요 시 공유드립니다.",
    10.5, LGRAY)
band(s, "핵심은 '사람을 줄이는 것'이 아니라, 같은 인력으로 더 빠르고·정확하게·많이 해내는 것입니다.")
footer(s, 13)

# =====================================================================
# SLIDE 14 — 앞으로의 계획 / 비전
# =====================================================================
s = slide()
header(s, KICK, "06", "앞으로의 계획 — 확장의 청사진",
       "더 개발하겠다는 게 아니라, 실제로 도입한다면 무엇이 필요한지를 제시합니다.")
# 좌: 확장 3단계
txt(s, 1.1, 2.0, 5.4, 0.3, "적용 범위는 이렇게 넓어집니다", 13, DARK, True)
stages = [
    ("지금 · 업무편람", "출처 기반 검색·답변  ✅ 구현·검증 완료"),
    ("다음 · 사내 문서 통합", "게시판·파일서버·PDF까지 같은 방식으로 확장"),
    ("비전 · 개인별 지식 허브", "메일·메신저·회의록까지 묶는 지식 허브로"),
]
y = 2.5
for i, (h, d) in enumerate(stages):
    card(s, 1.1, y, 5.3, 0.95, fill=(CREAM if i==0 else WHITE))
    txt(s, 1.35, y+0.16, 0.5, 0.5, "0"+str(i+1), 17, ORANGE, True)
    txt(s, 1.95, y+0.15, 4.3, 0.3, h, 12.5, DARK, True)
    txt(s, 1.95, y+0.5, 4.3, 0.35, d, 10, GRAY)
    y += 1.1
# 우: 실도입에 필요한 것
txt(s, 6.9, 2.0, 5.6, 0.3, "실도입에 필요한 것 (우리가 제시합니다)", 13, DARK, True)
needs = [
    ("사내 AI 환경", "AI를 돌릴 사내 GPU 서버 도입"),
    ("기업용 벡터 DB", "PoC용 경량 저장소 → 운영용 기업 DB"),
    ("국산 금융 LLM 전환", "데이터 주권·규제 대응을 위한 국산 모델"),
]
y = 2.5
for i, (h, d) in enumerate(needs):
    card(s, 6.9, y, 5.6, 0.95)
    txt(s, 7.15, y+0.15, 5.1, 0.3, h, 12.5, DARK, True)
    txt(s, 7.15, y+0.5, 5.1, 0.35, d, 10, GRAY)
    y += 1.1
txt(s, 6.9, 5.72, 5.6, 0.28, "타 증권사도 이미 같은 방향으로 가고 있습니다.", 10, ORANGE, True)
band(s, "이 프로젝트를 끌고 갈 역량이 우리에게 있다 — 그것을 오늘 보여드렸습니다.", t=6.12, h=0.58)
footer(s, 14)

# =====================================================================
# SLIDE 15 — Q&A
# =====================================================================
s = slide()
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = DARK; bg.line.fill.background(); bg.shadow.inherit=False
txt(s, 1.1, 2.7, 11.0, 1.2, "Q&A", 54, ORANGE, True)
txt(s, 1.15, 3.95, 11.0, 0.5, "Thank you", 22, WHITE, False)
txt(s, 1.15, 4.6, 11.0, 0.3, KICK, 11, FAINT)

# =====================================================================
# 부록 구분 슬라이드
# =====================================================================
s = slide()
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = DARK; bg.line.fill.background(); bg.shadow.inherit=False
txt(s, 1.1, 3.0, 11.0, 0.8, "부록 · 질문 대비 백업", 30, WHITE, True)
txt(s, 1.15, 3.9, 11.0, 0.4, "발표 본편에서는 넘어가고, 질문이 나올 때만 펼치는 근거 자료입니다.", 13, FAINT)

# =====================================================================
# 부록 A — 검색을 정확하게 만든 방법 (4신호)
# =====================================================================
s = slide()
header(s, KICK, "A", "부록 A · 답변을 정확하게 만든 방법",
       "기본 검색이 놓치는 것을, 네 가지 신호를 겹쳐 메웠습니다.")
sig = [
    ("제목·본문 함께 보기", "표제어 질문은 제목이, 풀어쓴 질문은 본문이 잡아냅니다."),
    ("같은 뜻 여러 번 묻기", "질문을 여러 표현으로 던져, 하나라도 명중하면 살립니다."),
    ("뜻으로 + 단어로 찾기", "의미 검색에 키워드 검색을 겹쳐 종목코드·고유명사를 잡습니다."),
    ("사내 용어 사전", "'예수금=고객예탁금'처럼 사내 용어를 사전 한 줄로 즉시 메웁니다."),
]
cw, gap, x0, y = 2.78, 0.24, 0.94, 2.15
for i, (h, d) in enumerate(sig):
    x = x0 + i*(cw+gap)
    card(s, x, y, cw, 2.9)
    txt(s, x+0.26, y+0.28, 0.6, 0.35, "0"+str(i+1), 16, ORANGE, True)
    txt(s, x+0.26, y+0.82, cw-0.52, 0.7, h, 13, DARK, True, line=1.1)
    txt(s, x+0.26, y+1.6, cw-0.52, 1.2, d, 10, GRAY, line=1.3)
band(s, "네 신호를 겹친 결과, 평가 문항에서 정답이 항상 상위 3건 안에 들어왔습니다.")
footer(s, "A")

# =====================================================================
# 부록 B — 정량 근거 수치 (물어볼 때만)
# =====================================================================
s = slide()
header(s, KICK, "B", "부록 B · 정량 근거 수치",
       "본편에서는 굳이 내세우지 않되, 물어보시면 바로 답할 수 있는 근거들입니다.")
stats = [
    ("업무편람 수집", "520p", "누락 0건"),
    ("적재 조각 수", "1,281", "의미 단위 색인"),
    ("검색 정확도", "상위 3건 내", "정답 항상 포함"),
    ("출처 부착률", "100%", "모든 답변에"),
    ("재처리 절감", "변경분만", "매일 자동 갱신"),
]
cw, gap, x0, y = 2.15, 0.28, 1.1, 2.3
for i, (h, v, d) in enumerate(stats):
    x = x0 + i*(cw+gap)
    card(s, x, y, cw, 2.1)
    txt(s, x+0.2, y+0.28, cw-0.4, 0.3, h, 11, GRAY, True, align=PP_ALIGN.CENTER)
    txt(s, x+0.2, y+0.72, cw-0.4, 0.55, v, 20, ORANGE, True, align=PP_ALIGN.CENTER)
    txt(s, x+0.2, y+1.45, cw-0.4, 0.4, d, 10, GRAY, align=PP_ALIGN.CENTER)
txt(s, 1.1, 4.7, 11.4, 0.6,
    "※ 답변 생성 정확도는 30문항 기준 측정값이며, 문항 수를 늘려 검증을 확대하고 있습니다. "
    "억원·ROI 등 비용 효과 수치는 산정 근거와 함께 별도 자료로 보관합니다.",
    10, LGRAY, line=1.3)
band(s, "수치는 '증명'을 위한 근거일 뿐, 오늘의 메시지는 '이만큼 만드는 팀'입니다.")
footer(s, "B")

# =====================================================================
# 부록 C — 타 금융사 RAG 현황
# =====================================================================
s = slide()
header(s, KICK, "C", "부록 C · 타 금융사 RAG 도입 현황",
       "검증 안 된 실험이 아니라, 업계가 이미 가고 있는 방향을 우리가 직접 구현했습니다.")
refs = [
    ("하나은행 H-GPT", "검색·AI 모드 + 규정 RAG + 출처 표시 (우리와 UX 사실상 동일)"),
    ("한국은행 BOKI", "완전 폐쇄망, 내부문서 140만건 (폐쇄망 선례의 최상위 근거)"),
    ("신한·미래에셋·키움·한투·IBK", "증권사 RAG 챗봇을 온프레미스로 도입 중 (국산 sLLM 포함)"),
    ("일본 MUFG·미즈호·SMBC", "사내 AI로 대규모 문서 검색·업무 시간 절감"),
]
y = 2.15
for i, (h, d) in enumerate(refs):
    card(s, 1.1, y, 11.4, 0.82, fill=(CREAM if i<2 else WHITE))
    txt(s, 1.35, y+0.14, 3.9, 0.55, h, 12.5, DARK, True, anchor=MSO_ANCHOR.MIDDLE, line=1.05)
    txt(s, 5.4, y+0.14, 6.9, 0.55, d, 11, GRAY, anchor=MSO_ANCHOR.MIDDLE, line=1.15)
    y += 0.95
band(s, "'검증된 방향을 벤더가 아니라 ITO팀이 직접 만들었다' — 그것이 차별점입니다.")
footer(s, "C")

out = "/home/user/2026_HICT_ChatBot/docs/presentation/증권 발표용/HICT_PPT_15min_신규.pptx"
prs.save(out)
print("saved:", out)
print("slides:", len(prs.slides.__iter__.__self__._sldIdLst))
